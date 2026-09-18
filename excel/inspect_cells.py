import openpyxl

wb = openpyxl.load_workbook("excel/AX_Portfolio_Manager_ROI_Simulator.xlsx")

def inspect_sheet(sheet_name, max_r=40, max_c=12):
    ws = wb[sheet_name]
    print(f"\n--- {sheet_name} ---")
    for r in range(1, max_r + 1):
        row_vals = []
        for c in range(1, max_c + 1):
            val = ws.cell(row=r, column=c).value
            if val is not None:
                row_vals.append(f"{openpyxl.utils.get_column_letter(c)}{r}: {str(val)[:25]}")
        if row_vals:
            print(" | ".join(row_vals[:6]))

inspect_sheet("04_TCO_8대비용군_산정", max_r=37, max_c=10)
inspect_sheet("05_편익_5대항목_산정", max_r=32, max_c=9)
inspect_sheet("06_ROI_현금흐름_시뮬레이터", max_r=26, max_c=10)
