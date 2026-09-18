import openpyxl

def verify_workbook():
    wb = openpyxl.load_workbook("excel/AX_Portfolio_Manager_ROI_Simulator.xlsx")
    print("Sheets in workbook:", wb.sheetnames)
    total_formulas = 0
    formula_samples = []

    for name in wb.sheetnames:
        ws = wb[name]
        formulas_in_sheet = 0
        for row in ws.iter_rows():
            for cell in row:
                val = str(cell.value)
                if val.startswith("="):
                    formulas_in_sheet += 1
                    total_formulas += 1
                    if len(formula_samples) < 15:
                        formula_samples.append(f"{name}!{cell.coordinate}: {val}")
        print(f"Sheet '{name}': {formulas_in_sheet} formulas, dimensions: {ws.dimensions}")

    print(f"\nTotal formulas in workbook: {total_formulas}")
    print("\nSample formulas:")
    for s in formula_samples:
        print(" ", s)

if __name__ == "__main__":
    verify_workbook()
